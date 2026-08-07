#!/usr/bin/env python3
"""Iki SAP ZANALYSIS_PATTERN ciktisini birlestirip en yogun 20 istasyonu raporlar."""
import collections
import os
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo

from parse import parse_file
from sources import source_files, HERE

SRC = source_files()          # [0] = yeni yillar, [1] = eski yillar
OUT = os.environ.get(
    'ZANALYSIS_OUT',
    os.path.join(os.path.dirname(HERE), '..', 'Istasyon_Bazli_Yuk_Ilk20.xlsx'))
OUT = os.path.abspath(OUT)

FONT = 'Arial'
NUM = '#,##0.00;-#,##0.00;-'
PCT = '0.0%;-0.0%;-'
HDR_FILL = PatternFill('solid', fgColor='1F3864')
SUB_FILL = PatternFill('solid', fgColor='2E5496')
TOT_FILL = PatternFill('solid', fgColor='FFF2CC')
ALT_FILL = PatternFill('solid', fgColor='F2F5FA')
TOP_FILL = PatternFill('solid', fgColor='E2EFDA')
THIN = Side(style='thin', color='B4B4B4')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def norm(s):
    """Turkce duyarli isim normalizasyonu (I/İ/ı/i ayrimi dahil)."""
    s = s.replace('İ', 'i').replace('I', 'ı').lower()
    s = unicodedata.normalize('NFD', s)
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    for a, b in (('ı', 'i'), ('ş', 's'), ('ğ', 'g'),
                 ('ç', 'c'), ('ö', 'o'), ('ü', 'u')):
        s = s.replace(a, b)
    return re.sub(r'[^a-z0-9]+', '', s)


# ---------------------------------------------------------------- veri okuma
parsed = [parse_file(p, lab) for p, lab in SRC]
SRCNAME = [lab for _, lab in SRC]
YEARS = sorted({y for d in parsed for y in d['years']})

# normalize edilmis anahtar -> gorunen ad (en cok tekrar eden yazim secilir)
name_votes = collections.defaultdict(collections.Counter)
for d in parsed:
    for n, y, f, v in d['records']:
        name_votes[norm(n)][n] += 1
display = {k: c.most_common(1)[0][0] for k, c in name_votes.items()}
# Islahiye icin dogru Turkce yazimi tercih et
for k, c in name_votes.items():
    if len(c) > 1:
        display[k] = max(c, key=lambda s: (s.startswith('İ'), c[s]))

merged = collections.defaultdict(float)   # (key, yil, akis) -> netton
for d in parsed:
    for n, y, f, v in d['records']:
        merged[(norm(n), y, f)] += v

keys = sorted(name_votes)
tot_of = {k: sum(merged.get((k, y, f), 0.0)
                 for y in YEARS for f in ('Giden', 'Gelen')) for k in keys}
order = sorted(keys, key=lambda k: (-tot_of[k], norm(display[k])))
TOP = order[:20]

GRAND = sum(tot_of.values())
print(f"birlesik istasyon: {len(keys)}  yillar: {YEARS[0]}-{YEARS[-1]}")
print(f"genel toplam (Giden+Gelen): {GRAND:,.2f}")

# --------------------------------------------------------------- yardimcilar
def style_header(ws, row, first, last, text_rot=False, fill=HDR_FILL):
    for c in range(first, last + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = Font(name=FONT, size=9, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center',
                                   wrap_text=True,
                                   textRotation=90 if text_rot else 0)
        cell.border = BORDER


def put(ws, r, c, val, *, num=None, bold=False, fill=None,
        align=None, size=9, color='000000'):
    cell = ws.cell(row=r, column=c, value=val)
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    cell.border = BORDER
    if num:
        cell.number_format = num
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = Alignment(horizontal=align, vertical='center')
    return cell


def setup_print(ws, *, title_cols, title_rows, area, landscape=True,
                fit_w=1, fit_h=0):
    """PDF / yazdirma duzeni.

    ONEMLI: freeze_panes (Bölmeleri Dondur) yalnizca ekranda gecerlidir ve
    yazdirmaya/PDF'e hicbir etkisi yoktur. Her sayfada ilk sutunun (ve baslik
    satirinin) yinelenmesi icin print_title_cols / print_title_rows gerekir —
    Excel arayuzunde: Sayfa Düzeni > Yazdırma Başlıkları >
    'Soldan yinelenecek sütunlar' ve 'Üstte yinelenecek satırlar'.
    """
    ws.print_title_cols = title_cols      # her sayfada solda yinelenir
    ws.print_title_rows = title_rows      # her sayfada ustte yinelenir
    ws.print_area = area
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = fit_w      # genisligi N sayfaya sigdir
    ws.page_setup.fitToHeight = fit_h     # 0 = yukseklikte serbest
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.5
    ws.page_margins.header = ws.page_margins.footer = 0.25
    ws.oddHeader.left.text = '&"Arial,Bold"&10' + ws.title
    ws.oddHeader.right.text = '&"Arial"&8İstasyon Bazlı Yük Taşımaları'
    ws.oddFooter.left.text = '&"Arial"&8Kaynak: SAP ZANALYSIS_PATTERN (birleşik)'
    ws.oddFooter.right.text = '&"Arial"&8Sayfa &P / &N'


wb = Workbook()

# ============================================================ 3. VERI SAYFASI
# Once master veri sayfasi kurulur; ozet sayfalari buraya formulle baglanir.
wsd = wb.create_sheet('Tüm İstasyonlar (Birleşik)')
wsd.sheet_view.showGridLines = False

put(wsd, 1, 1, 'Birleşik Ham Veri — İstasyon Bazlı Yük Taşımaları (Netton, TON)',
    bold=True, size=12).border = Border()
put(wsd, 2, 1, f'Kaynak: {SRC[1][1]} ({parsed[1]["years"][0]}-{parsed[1]["years"][-1]}) '
               f'+ {SRC[0][1]} ({parsed[0]["years"][0]}-{parsed[0]["years"][-1]})',
    size=8, color='595959').border = Border()

HR1, HR2 = 4, 5      # baslik satirlari
DR0 = 6              # ilk veri satiri
put(wsd, HR1, 1, 'Sıra')
put(wsd, HR1, 2, 'İstasyon')
wsd.merge_cells(start_row=HR1, start_column=1, end_row=HR2, end_column=1)
wsd.merge_cells(start_row=HR1, start_column=2, end_row=HR2, end_column=2)

col_of = {}
c = 3
for y in YEARS:
    wsd.merge_cells(start_row=HR1, start_column=c, end_row=HR1, end_column=c + 1)
    put(wsd, HR1, c, y)
    put(wsd, HR2, c, 'Giden')
    put(wsd, HR2, c + 1, 'Gelen')
    col_of[(y, 'Giden')], col_of[(y, 'Gelen')] = c, c + 1
    c += 2
C_GID, C_GEL, C_TOP = c, c + 1, c + 2
wsd.merge_cells(start_row=HR1, start_column=C_GID, end_row=HR1, end_column=C_TOP)
put(wsd, HR1, C_GID, f'TOPLAM {YEARS[0]}-{YEARS[-1]}')
put(wsd, HR2, C_GID, 'Giden')
put(wsd, HR2, C_GEL, 'Gelen')
put(wsd, HR2, C_TOP, 'Genel Toplam')
LASTC = C_TOP
style_header(wsd, HR1, 1, LASTC)
style_header(wsd, HR2, 1, LASTC, fill=SUB_FILL)
wsd.row_dimensions[HR1].height = 18
wsd.row_dimensions[HR2].height = 16

row_of = {}
for i, k in enumerate(order):
    r = DR0 + i
    row_of[k] = r
    shade = TOP_FILL if i < 20 else (ALT_FILL if i % 2 else None)
    put(wsd, r, 1, i + 1, num='0', align='center', fill=shade)
    put(wsd, r, 2, display[k], bold=i < 20, fill=shade)
    for y in YEARS:
        for f in ('Giden', 'Gelen'):
            v = merged.get((k, y, f))
            put(wsd, r, col_of[(y, f)], v, num=NUM, fill=shade)
    gi = [get_column_letter(col_of[(y, 'Giden')]) for y in YEARS]
    ge = [get_column_letter(col_of[(y, 'Gelen')]) for y in YEARS]
    put(wsd, r, C_GID, '=SUM(' + ','.join(f'{x}{r}' for x in gi) + ')',
        num=NUM, bold=True, fill=shade)
    put(wsd, r, C_GEL, '=SUM(' + ','.join(f'{x}{r}' for x in ge) + ')',
        num=NUM, bold=True, fill=shade)
    put(wsd, r, C_TOP,
        f'={get_column_letter(C_GID)}{r}+{get_column_letter(C_GEL)}{r}',
        num=NUM, bold=True, fill=shade)

DRN = DR0 + len(order) - 1
TR = DRN + 1
put(wsd, TR, 2, 'GENEL TOPLAM', bold=True, fill=TOT_FILL)
put(wsd, TR, 1, None, fill=TOT_FILL)
for cc in range(3, LASTC + 1):
    L = get_column_letter(cc)
    put(wsd, TR, cc, f'=SUM({L}{DR0}:{L}{DRN})', num=NUM, bold=True, fill=TOT_FILL)

wsd.column_dimensions['A'].width = 6
wsd.column_dimensions['B'].width = 30
for cc in range(3, LASTC + 1):
    wsd.column_dimensions[get_column_letter(cc)].width = 15
wsd.freeze_panes = f'C{DR0}'
wsd.auto_filter.ref = f'A{HR2}:{get_column_letter(LASTC)}{DRN}'

DSH = "'Tüm İstasyonlar (Birleşik)'"
GT_REF = f'{DSH}!${get_column_letter(C_TOP)}${TR}'

# ============================================================ 1. OZET SAYFASI
ws = wb.create_sheet('İlk 20 İstasyon', 0)
ws.sheet_view.showGridLines = False

put(ws, 1, 1, 'EN FAZLA YÜK GELİP GİDEN İLK 20 İSTASYON', bold=True,
    size=14, color='1F3864').border = Border()
put(ws, 2, 1, f'İstasyon Bazlı Yük Taşımaları — {YEARS[0]}-{YEARS[-1]} birleşik, '
              f'Netton (TON), Giden + Gelen toplamına göre sıralı',
    size=9, color='595959').border = Border()
# dosya bazli alt toplamlar icin yil gruplari
YB = parsed[1]['years']          # eski dosya
YA = parsed[0]['years']          # yeni dosya
SPAN_B = f'{YB[0]}-{YB[-1]}'
SPAN_A = f'{YA[0]}-{YA[-1]}'

put(ws, 3, 1, f'Kaynak: 2 adet SAP ZANALYSIS_PATTERN çıktısı birleştirilmiştir '
              f'({SRCNAME[1]} → {SPAN_B} ve {SRCNAME[0]} → {SPAN_A}). '
              f'{YEARS[-1]} yılı kısmi (yıl tamamlanmamıştır).',
    size=8, color='808080').border = Border()

H = 5
cols = ['Sıra', 'İstasyon', 'Giden (Netton)', 'Gelen (Netton)',
        'TOPLAM (Netton)', 'Pay %', 'Kümülatif %',
        f'{SPAN_B} Toplam', f'{SPAN_A} Toplam']
for j, t in enumerate(cols, start=1):
    put(ws, H, j, t)
style_header(ws, H, 1, len(cols))
ws.row_dimensions[H].height = 30


def span_sum(k, years):
    r = row_of[k]
    parts = []
    for y in years:
        parts.append(f'{DSH}!{get_column_letter(col_of[(y, "Giden")])}{r}')
        parts.append(f'{DSH}!{get_column_letter(col_of[(y, "Gelen")])}{r}')
    return '=SUM(' + ','.join(parts) + ')'


R0 = H + 1
for i, k in enumerate(TOP):
    r = R0 + i
    src = row_of[k]
    fill = ALT_FILL if i % 2 else None
    put(ws, r, 1, i + 1, num='0', align='center', fill=fill)
    put(ws, r, 2, display[k], bold=True, fill=fill)
    put(ws, r, 3, f'={DSH}!{get_column_letter(C_GID)}{src}', num=NUM, fill=fill)
    put(ws, r, 4, f'={DSH}!{get_column_letter(C_GEL)}{src}', num=NUM, fill=fill)
    put(ws, r, 5, f'=C{r}+D{r}', num=NUM, bold=True, fill=fill)
    put(ws, r, 6, f'=IFERROR(E{r}/{GT_REF},0)', num=PCT, fill=fill)
    put(ws, r, 7, f'=IFERROR(SUM($E${R0}:E{r})/{GT_REF},0)', num=PCT, fill=fill)
    put(ws, r, 8, span_sum(k, YB), num=NUM, fill=fill)
    put(ws, r, 9, span_sum(k, YA), num=NUM, fill=fill)

RN = R0 + len(TOP) - 1
r = RN + 1
put(ws, r, 2, 'İLK 20 TOPLAMI', bold=True, fill=TOT_FILL)
put(ws, r, 1, None, fill=TOT_FILL)
for cc, L in ((3, 'C'), (4, 'D'), (5, 'E'), (8, 'H'), (9, 'I')):
    put(ws, r, cc, f'=SUM({L}{R0}:{L}{RN})', num=NUM, bold=True, fill=TOT_FILL)
put(ws, r, 6, f'=IFERROR(E{r}/{GT_REF},0)', num=PCT, bold=True, fill=TOT_FILL)
put(ws, r, 7, None, fill=TOT_FILL)

r += 1
put(ws, r, 2, 'TÜM İSTASYONLAR (GENEL TOPLAM)', bold=True, fill=TOT_FILL)
put(ws, r, 1, None, fill=TOT_FILL)
put(ws, r, 3, f'={DSH}!{get_column_letter(C_GID)}{TR}', num=NUM, bold=True, fill=TOT_FILL)
put(ws, r, 4, f'={DSH}!{get_column_letter(C_GEL)}{TR}', num=NUM, bold=True, fill=TOT_FILL)
put(ws, r, 5, f'={GT_REF}', num=NUM, bold=True, fill=TOT_FILL)
put(ws, r, 6, f'=IFERROR(E{r}/{GT_REF},0)', num=PCT, bold=True, fill=TOT_FILL)
put(ws, r, 7, None, fill=TOT_FILL)
gb = '+'.join(f'{DSH}!{get_column_letter(col_of[(y, f)])}{TR}'
              for y in YB for f in ('Giden', 'Gelen'))
ga = '+'.join(f'{DSH}!{get_column_letter(col_of[(y, f)])}{TR}'
              for y in YA for f in ('Giden', 'Gelen'))
put(ws, r, 8, '=' + gb, num=NUM, bold=True, fill=TOT_FILL)
put(ws, r, 9, '=' + ga, num=NUM, bold=True, fill=TOT_FILL)
LAST_SUM_ROW = r

ws.cell(row=H, column=5).comment = Comment(
    'TOPLAM = Giden + Gelen (Netton, TON). Bir sevkiyat çıkış istasyonunda '
    '"Giden", varış istasyonunda "Gelen" olarak sayıldığından bu ölçü '
    'istasyonun toplam yük trafiğini gösterir; ülke geneli taşınan tonajı değil.',
    'Analiz', width=380, height=110)
ws.cell(row=H, column=6).comment = Comment(
    'Pay % = İstasyon Toplamı / Tüm istasyonların Genel Toplamı '
    '(Giden+Gelen bazında).', 'Analiz', width=320, height=70)

for col, w in zip('ABCDEFGHI', (6, 28, 17, 17, 18, 9, 12, 18, 18)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = f'C{R0}'                       # ekran: Sıra + İstasyon sabit
setup_print(ws, title_cols='$A:$B', title_rows=f'${H}:${H}',
            area=f'A1:I{LAST_SUM_ROW}')          # grafik ayri sayfada kalir

chart = BarChart()
chart.type = 'bar'
chart.style = 10
chart.title = f'İlk 20 İstasyon — Toplam Yük (Netton, {YEARS[0]}-{YEARS[-1]})'
chart.y_axis.title = 'Netton (TON)'
chart.add_data(Reference(ws, min_col=5, min_row=H, max_row=RN), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=2, min_row=R0, max_row=RN))
chart.height, chart.width = 12, 22
chart.gapWidth = 40
ws.add_chart(chart, f'K{H}')

# ================================================== 2. YIL BAZINDA ILK 20
wsy = wb.create_sheet('İlk 20 - Yıl Bazında', 1)
wsy.sheet_view.showGridLines = False
put(wsy, 1, 1, 'İlk 20 İstasyon — Yıl Bazında Yük (Netton, TON)',
    bold=True, size=12, color='1F3864').border = Border()
put(wsy, 2, 1, f'{SPAN_B} verisi {SRCNAME[1]}, {SPAN_A} verisi {SRCNAME[0]} '
               f'dosyasından gelmektedir. {YEARS[-1]} kısmi yıldır.',
    size=8, color='808080').border = Border()

YH1, YH2, YD0 = 4, 5, 6
put(wsy, YH1, 1, 'Sıra')
put(wsy, YH1, 2, 'İstasyon')
wsy.merge_cells(start_row=YH1, start_column=1, end_row=YH2, end_column=1)
wsy.merge_cells(start_row=YH1, start_column=2, end_row=YH2, end_column=2)
ycol = {}
c = 3
for y in YEARS:
    wsy.merge_cells(start_row=YH1, start_column=c, end_row=YH1, end_column=c + 2)
    put(wsy, YH1, c, y + (' *' if y == YEARS[-1] else ''))
    for k2, t in enumerate(('Giden', 'Gelen', 'Toplam')):
        put(wsy, YH2, c + k2, t)
    ycol[y] = c
    c += 3
YT = c
put(wsy, YH1, YT, 'GENEL\nTOPLAM')
wsy.merge_cells(start_row=YH1, start_column=YT, end_row=YH2, end_column=YT)
style_header(wsy, YH1, 1, YT)
style_header(wsy, YH2, 1, YT, fill=SUB_FILL)
wsy.row_dimensions[YH1].height = 20
wsy.row_dimensions[YH2].height = 16

for i, k in enumerate(TOP):
    r = YD0 + i
    src = row_of[k]
    fill = ALT_FILL if i % 2 else None
    put(wsy, r, 1, i + 1, num='0', align='center', fill=fill)
    put(wsy, r, 2, display[k], bold=True, fill=fill)
    for y in YEARS:
        c = ycol[y]
        put(wsy, r, c,
            f'={DSH}!{get_column_letter(col_of[(y, "Giden")])}{src}', num=NUM, fill=fill)
        put(wsy, r, c + 1,
            f'={DSH}!{get_column_letter(col_of[(y, "Gelen")])}{src}', num=NUM, fill=fill)
        put(wsy, r, c + 2,
            f'={get_column_letter(c)}{r}+{get_column_letter(c+1)}{r}',
            num=NUM, bold=True, fill=fill)
    put(wsy, r, YT,
        '=' + '+'.join(f'{get_column_letter(ycol[y]+2)}{r}' for y in YEARS),
        num=NUM, bold=True, fill=fill)

YN = YD0 + len(TOP) - 1
r = YN + 1
put(wsy, r, 2, 'İLK 20 TOPLAMI', bold=True, fill=TOT_FILL)
put(wsy, r, 1, None, fill=TOT_FILL)
for cc in range(3, YT + 1):
    L = get_column_letter(cc)
    put(wsy, r, cc, f'=SUM({L}{YD0}:{L}{YN})', num=NUM, bold=True, fill=TOT_FILL)
r += 1
put(wsy, r, 2, 'TÜM İSTASYONLAR', bold=True, fill=TOT_FILL)
put(wsy, r, 1, None, fill=TOT_FILL)
for y in YEARS:
    c = ycol[y]
    put(wsy, r, c, f'={DSH}!{get_column_letter(col_of[(y, "Giden")])}{TR}',
        num=NUM, bold=True, fill=TOT_FILL)
    put(wsy, r, c + 1, f'={DSH}!{get_column_letter(col_of[(y, "Gelen")])}{TR}',
        num=NUM, bold=True, fill=TOT_FILL)
    put(wsy, r, c + 2, f'={get_column_letter(c)}{r}+{get_column_letter(c+1)}{r}',
        num=NUM, bold=True, fill=TOT_FILL)
put(wsy, r, YT, '=' + '+'.join(f'{get_column_letter(ycol[y]+2)}{r}' for y in YEARS),
    num=NUM, bold=True, fill=TOT_FILL)
r += 1
put(wsy, r, 2, '* 2026 yılı kısmi veridir.', size=8, color='808080').border = Border()

wsy.column_dimensions['A'].width = 6
wsy.column_dimensions['B'].width = 28
for cc in range(3, YT + 1):
    wsy.column_dimensions[get_column_letter(cc)].width = 14
wsy.freeze_panes = f'C{YD0}'

# ============================================================= 4. YONTEM
wsm = wb.create_sheet('Yöntem ve Notlar')
wsm.sheet_view.showGridLines = False
wsm.column_dimensions['A'].width = 32
wsm.column_dimensions['B'].width = 105

put(wsm, 1, 1, 'YÖNTEM VE NOTLAR', bold=True, size=13, color='1F3864').border = Border()
notes = [
    ('Amaç', 'İki ayrı SAP ZANALYSIS_PATTERN çıktısındaki istasyon bazlı yük '
             'taşıma verilerini birleştirip, gelen + giden yük toplamı en yüksek '
             '20 istasyonu sıralamak.'),
    ('Kaynak dosya 1', f'{SRC[1][1]} — "İstasyon Bazlı Yük Taşımaları", '
                       f'{parsed[1]["years"][0]}-{parsed[1]["years"][-1]}, '
                       f'{len({r[0] for r in parsed[1]["records"]})} istasyon.'),
    ('Kaynak dosya 2', f'{SRC[0][1]} — "İstasyona Göre Yük Taşımaları", '
                       f'{parsed[0]["years"][0]}-{parsed[0]["years"][-1]}, '
                       f'{len({r[0] for r in parsed[0]["records"]})} istasyon.'),
    ('Birleştirme', 'İki dosyanın yıl aralıkları çakışmadığı için (2013-2016 ve '
                    '2017-2026) mükerrer sayım riski yoktur. Veriler istasyon adı '
                    'üzerinden eşleştirilerek yan yana birleştirilmiştir. '
                    f'Birleşik liste: {len(keys)} istasyon, {YEARS[0]}-{YEARS[-1]}.'),
    ('Ölçü birimi', 'Netton (TON). Kaynak dosyalardaki "Giden" ve "Gelen" '
                    'sütunları kullanılmıştır.'),
    ('Sıralama ölçütü', 'TOPLAM = Giden + Gelen. Bu, istasyonun toplam yük '
                        'trafiğidir. Bir sevkiyat çıkış istasyonunda "Giden", '
                        'varış istasyonunda "Gelen" olarak sayıldığından, tüm '
                        'istasyon toplamları ülke genelinde taşınan net tonajın '
                        'yaklaşık iki katına denk gelir.'),
    ('İsim eşleştirme', 'İstasyon adları Türkçe büyük/küçük harf kuralları (I/İ/ı/i) '
                        've aksan farkları normalize edilerek eşleştirilmiştir. '
                        'Tespit edilen tek yazım farkı: "Islahiye" (2017-2026 dosyası) '
                        'ile "İslahiye" (2013-2016 dosyası) aynı istasyon kabul edilip '
                        'birleştirilmiştir.'),
    ('Mükerrer satır', '2013-2016 dosyasında "Azot" adıyla iki ayrı satır bulunmaktadır '
                       '(442.056,46 ve 599,11 Netton). Aynı istasyon adı altında '
                       'toplanmıştır.'),
    ('Ayrı tutulan adlar', '"Mersin" / "Mersin Liman" ve "Haydarpaşa" / "Haydarpaşa Liman" '
                           'kaynak dosyalarda ayrı kayıtlar olduğu için ayrı '
                           'istasyon olarak bırakılmıştır.'),
    ('2026 verisi', '2026 yılı kısmi veridir (dosya 07.08.2026 tarihinde alınmıştır), '
                    'tam yıl ile karşılaştırılırken dikkate alınmalıdır.'),
    ('Doğrulama', 'Ayrıştırma, kaynak dosyaların kendi ara toplamlarıyla karşılaştırılarak '
                  'doğrulanmıştır: her satırda "Sonuç = Giden + Gelen" ve '
                  '"Toplam sonuç = yıllık Sonuç toplamı" kontrolleri 0 hata vermiştir. '
                  '2013-2016 dosyasının kendi "Toplam sonuç" satırı 207.025.948,53 Netton '
                  'olup hesaplanan değerle birebir örtüşmektedir. '
                  '(2017-2026 dosyasında genel toplam satırı bulunmadığından satır ve '
                  'sütun bazlı iç tutarlılık kontrolü uygulanmıştır.)'),
    ('Sayfalar', '"İlk 20 İstasyon": sıralı özet ve grafik. '
                 '"İlk 20 - Yıl Bazında": ilk 20 istasyonun 2013-2026 yıllık dökümü. '
                 '"Tüm İstasyonlar (Birleşik)": birleştirilmiş ham veri; diğer '
                 'sayfalardaki tüm değerler bu sayfaya formülle bağlıdır.'),
]
r = 3
for k2, v in notes:
    a = put(wsm, r, 1, k2, bold=True)
    a.alignment = Alignment(vertical='top')
    b = put(wsm, r, 2, v)
    b.alignment = Alignment(wrap_text=True, vertical='top')
    wsm.row_dimensions[r].height = max(15, 13 * (len(v) // 100 + 1))
    r += 1

del wb['Sheet']
wb.move_sheet('Tüm İstasyonlar (Birleşik)', offset=0)
wb._sheets = [wb['İlk 20 İstasyon'], wb['İlk 20 - Yıl Bazında'],
              wb['Tüm İstasyonlar (Birleşik)'], wb['Yöntem ve Notlar']]
wb.active = 0
wb.save(OUT)
print('yazildi:', OUT)

print('\n== İLK 20 (python ile hesaplanan referans) ==')
for i, k in enumerate(TOP, 1):
    gi = sum(merged.get((k, y, 'Giden'), 0.0) for y in YEARS)
    ge = sum(merged.get((k, y, 'Gelen'), 0.0) for y in YEARS)
    print(f'{i:>3}. {display[k]:<26} Giden {gi:>16,.2f}  Gelen {ge:>16,.2f}  '
          f'Toplam {tot_of[k]:>16,.2f}  Pay {tot_of[k]/GRAND:6.2%}')
print(f'{"":>4} {"İLK 20 TOPLAMI":<26} {"":>16} {"":>18} '
      f'Toplam {sum(tot_of[k] for k in TOP):>16,.2f}  '
      f'Pay {sum(tot_of[k] for k in TOP)/GRAND:6.2%}')
print(f'{"":>4} {"GENEL TOPLAM":<26} {"":>16} {"":>18} Toplam {GRAND:>16,.2f}')
